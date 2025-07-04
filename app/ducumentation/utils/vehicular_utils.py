import re
from docx import Document

def replace_placeholders(doc, data):
    """
    Replace placeholders in the document with actual data.
    """
    for paragraph in doc.paragraphs:
        for placeholder, value in data.items():
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, value)
    return doc


def remove_placeholders(doc):
    """
    Remove unused placeholders from the document.
    """
    for paragraph in doc.paragraphs:
        paragraph.text = re.sub(r"\[E\.[^\]]+\]", "", paragraph.text)  # Remove placeholders like [E.PLACEHOLDER]
    return doc

def format_data(data, count, gender):
    """
    Format data for singular/plural and gender-specific text.
    """
    if count > 1:
        if gender == "F":
            data["CALIDAD"] = "Compradoras"
            data["EL"] = "Las"
        else:
            data["CALIDAD"] = "Compradores"
            data["EL"] = "Los"
    else:
        if gender == "F":
            data["CALIDAD"] = "Compradora"
            data["EL"] = "La"
        else:
            data["CALIDAD"] = "Comprador"
            data["EL"] = "El"
    return data

def add_default_values(data, placeholders):
    """
    Add default values for missing placeholders.
    """
    for placeholder in placeholders:
        if placeholder not in data:
            data[placeholder] = f"[{placeholder}]"
    return data